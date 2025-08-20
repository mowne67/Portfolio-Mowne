import pandas as pd
import numpy as np
import re
from openpyxl import load_workbook
from openpyxl.styles import Font

def process_excel(file_path):
    # --- Function to clean currency strings to floats ---
    def parse_currency(val):
        if pd.isna(val) or str(val).strip().upper() == "NA":
            return np.nan
        cleaned = re.sub(r"[^\d\.\-]", "", str(val))
        return float(cleaned) if cleaned else np.nan

    # --- Load workbook for sheet names ---
    wb = load_workbook(file_path)
    sheet_names = [s for s in wb.sheetnames if s.startswith("App")]
    wb.close()

    # --- Process each App sheet ---
    for sheet_name in sheet_names:
        df = pd.read_excel(file_path, sheet_name=sheet_name)

        if not {"fma", "calculation"}.issubset(df.columns):
            continue  # Skip if required columns don't exist

        # Clean numeric versions
        df["fma_num"] = df["fma"].apply(parse_currency)
        df["calc_num"] = df["calculation"].apply(parse_currency)

        # Apply decision logic
        def decision_logic(row):
            if pd.isna(row["fma_num"]) or pd.isna(row["calc_num"]):
                return ""
            if row["fma_num"] > row["calc_num"]:
                return row["fma"]
            elif row["fma_num"] != row["calc_num"]:
                return row["fma"]  # Will be colored red later
            else:
                return row["calculation"]

        df["decision"] = df.apply(decision_logic, axis=1)

        # Drop helper columns
        df.drop(["fma_num", "calc_num"], axis=1, inplace=True)

        # Write back to Excel (overwrite only this sheet)
        with pd.ExcelWriter(file_path, mode="a", engine="openpyxl", if_sheet_exists="replace") as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False)

        # Apply red font formatting where needed
        wb = load_workbook(file_path)
        ws = wb[sheet_name]

        # Get column indices
        headers = {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}
        decision_col = headers.get("decision")
        fma_col = headers.get("fma")
        calc_col = headers.get("calculation")

        for row_idx in range(2, ws.max_row + 1):
            fma_val = parse_currency(ws.cell(row=row_idx, column=fma_col).value)
            calc_val = parse_currency(ws.cell(row=row_idx, column=calc_col).value)
            decision_val = ws.cell(row=row_idx, column=decision_col).value

            if fma_val != calc_val and decision_val != "" and fma_val <= calc_val:
                ws.cell(row=row_idx, column=decision_col).font = Font(color="FF0000")  # red text

        wb.save(file_path)

    print(f"✅ Processing completed for sheets starting with 'App' in {file_path}")
